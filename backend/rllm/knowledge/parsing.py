"""
Document parsing → structure-preserving text blocks.

Every format is reduced to a list of :class:`Block` (a paragraph-ish unit
with its page number and section heading). Keeping page/section lets
citations say *where* in a document the evidence came from.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Dict, List


@dataclass
class Block:
    text: str
    page: int = 0          # 1-based for paged formats, 0 otherwise
    section: str = ""      # nearest heading


class ParseError(ValueError):
    pass


SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".md", ".markdown", ".txt", ".html", ".htm")


def _clean(text: str) -> str:
    text = text.replace(" ", " ").replace("​", "")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


# ── PDF ─────────────────────────────────────────────────────────────────────
def _parse_pdf(data: bytes) -> List[Block]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"Could not open PDF: {e}")
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ParseError("PDF is password-protected")
    blocks: List[Block] = []
    for pno, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        raw = re.sub(r"(\w)-\n(\w)", r"\1\2", raw)          # de-hyphenate line breaks
        for para in re.split(r"\n\s*\n", raw):
            para = _clean(para.replace("\n", " "))
            if para:
                blocks.append(Block(para, page=pno))
    return blocks


# ── DOCX ────────────────────────────────────────────────────────────────────
def _parse_docx(data: bytes) -> List[Block]:
    try:
        import docx  # python-docx
    except ImportError:
        raise ParseError("DOCX support needs python-docx:  pip install python-docx")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"Could not open DOCX: {e}")
    blocks: List[Block] = []
    section = ""
    for para in document.paragraphs:
        text = _clean(para.text)
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            section = text
            continue
        blocks.append(Block(text, section=section))
    for table in document.tables:
        for row in table.rows:
            cells = [_clean(c.text) for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                blocks.append(Block(line, section=section or "Table"))
    return blocks


# ── Markdown / text ────────────────────────────────────────────────────────
_MD_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")


def _strip_md_inline(s: str) -> str:
    s = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", s)      # images → alt text
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)       # links → text
    s = re.sub(r"`([^`]+)`", r"\1", s)
    s = re.sub(r"(\*\*|__)(.+?)\1", r"\2", s)
    s = re.sub(r"(?<!\w)[*_](.+?)[*_](?!\w)", r"\1", s)
    s = re.sub(r"<[^>]+>", "", s)
    return s


def _parse_markdown_text(text: str) -> List[Block]:
    blocks: List[Block] = []
    section = ""
    buf: List[str] = []
    in_code = False

    def flush():
        if buf:
            para = _clean(" ".join(buf))
            if para:
                blocks.append(Block(para, section=section))
            buf.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            buf.append(line.strip())
            continue
        m = _MD_HEADING.match(line)
        if m:
            flush()
            section = _strip_md_inline(m.group(2)).strip()
            continue
        stripped = line.strip()
        if not stripped or re.fullmatch(r"[-*_=]{3,}", stripped):
            flush()
            continue
        stripped = re.sub(r"^(?:[-*+]|\d+[.)])\s+", "", stripped)   # list markers
        stripped = re.sub(r"^>\s?", "", stripped)                   # blockquote
        if stripped.startswith("|"):
            stripped = " | ".join(c.strip() for c in stripped.strip("|").split("|")
                                  if c.strip() and not re.fullmatch(r":?-+:?", c.strip()))
            if not stripped:
                continue
        buf.append(_strip_md_inline(stripped))
    flush()
    return blocks


def _parse_markdown(data: bytes) -> List[Block]:
    return _parse_markdown_text(data.decode("utf-8", errors="replace"))


def _parse_txt(data: bytes) -> List[Block]:
    text = data.decode("utf-8", errors="replace")
    return [Block(_clean(p.replace("\n", " "))) for p in re.split(r"\n\s*\n", text) if _clean(p)]


# ── HTML ────────────────────────────────────────────────────────────────────
_HTML_TAGS = ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "blockquote", "td", "th", "dd", "dt", "figcaption"]


def _parse_html(data: bytes) -> List[Block]:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data.decode("utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "noscript", "head", "nav", "footer", "svg"]):
        tag.decompose()
    blocks: List[Block] = []
    section = ""
    for el in soup.find_all(_HTML_TAGS):
        if el.find(_HTML_TAGS):          # only leaf-level text containers
            continue
        text = _clean(el.get_text(" "))
        if not text:
            continue
        if el.name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            section = text
            continue
        blocks.append(Block(text, section=section))
    if not blocks:  # unstructured page — fall back to all visible text
        text = _clean(soup.get_text(" "))
        if text:
            blocks.append(Block(text))
    return blocks


_PARSERS: Dict[str, Callable[[bytes], List[Block]]] = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".md": _parse_markdown,
    ".markdown": _parse_markdown,
    ".txt": _parse_txt,
    ".html": _parse_html,
    ".htm": _parse_html,
}


def file_type_for(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    return {".htm": "html", ".markdown": "md"}.get(ext, ext.lstrip("."))


def parse_document(data: bytes, filename: str) -> List[Block]:
    ext = Path(filename).suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ParseError(f"Unsupported file type '{ext or filename}'. "
                         f"Supported: {', '.join(SUPPORTED_EXTENSIONS)}")
    blocks = parser(data)
    if not blocks:
        raise ParseError("No extractable text found (scanned PDFs need OCR first)")
    return blocks
